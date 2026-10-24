import streamlit as st
import pandas as pd
import os
import zipfile
import re
import base64
import gdown
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

FAVICON_OFICIAL = "https://cdn.prod.website-files.com/69cb021343dbf125cabfa994/69cb021343dbf125cabfaa24_favicon-suzano.png"
st.set_page_config(page_title="Suzano | Trade Marketing", page_icon=FAVICON_OFICIAL, layout="wide")

LINK_DRIVE_ZIP = "https://drive.google.com/file/d/18GBHvpYRNLVLvDE8EdUz2QQXOMPpnbPd/view?usp=sharing"
LINK_DRIVE_CAPA = "https://docs.google.com/presentation/d/1Q4DfabWYLGPORt1t-N4zF87hA9yh1m7d/export?format=pptx"

PASTA_FOTOS = "fotos_temp"
ARQUIVO_ZIP = "Banco_de_Fotos_Suzano.zip"
ARQUIVO_CAPA = "Base_Capa.pptx"

@st.cache_resource(show_spinner="A carregar catálogo visual e templates da Suzano...")
def preparar_arquivos_nuvem():
    if not os.path.exists(PASTA_FOTOS):
        if not os.path.exists(ARQUIVO_ZIP):
            try:
                gdown.download(LINK_DRIVE_ZIP, ARQUIVO_ZIP, quiet=False, fuzzy=True)
            except Exception:
                gdown.download("https://drive.google.com/uc?id=18GBHvpYRNLVLvDE8EdUz2QQXOMPpnbPd", ARQUIVO_ZIP, quiet=False)
        
        if os.path.exists(ARQUIVO_ZIP):
            with zipfile.ZipFile(ARQUIVO_ZIP, 'r') as z:
                z.extractall(PASTA_FOTOS)
                
    if not os.path.exists(ARQUIVO_CAPA):
        try:
            gdown.download(LINK_DRIVE_CAPA, ARQUIVO_CAPA, quiet=False, fuzzy=True)
        except Exception:
            gdown.download("https://docs.google.com/presentation/d/1Q4DfabWYLGPORt1t-N4zF87hA9yh1m7d/export?format=pptx", ARQUIVO_CAPA, quiet=False)

preparar_arquivos_nuvem()

def achar_logo(marca):
    m = marca.upper().strip()
    if not os.path.exists(PASTA_FOTOS): return None
    
    for root, _, files in os.walk(PASTA_FOTOS):
        for f in files:
            f_up = f.upper()
            if m in f_up:
                if m == "NEVE":
                    if not any(sub in f_up for sub in ["FOLHA", "DUPLA", "TRIPLA", "WIPES", "QUADRUPLA", "ROLO", "TOQUE", "PACOTE"]):
                        return os.path.join(root, f)
                elif f"LOGO_{m}" in f_up or m in f_up:
                    return os.path.join(root, f)
                    
    for root, _, files in os.walk(PASTA_FOTOS):
        for f in files:
            f_up = f.upper()
            if m in f_up and not any(c.isdigit() for c in f_up):
                return os.path.join(root, f)
    return None

def achar_imagem(mat_code):
    codigo_busca = str(mat_code).split("/")[0].strip().upper()
    if not os.path.exists(PASTA_FOTOS): return None
    for root, _, files in os.walk(PASTA_FOTOS):
        for f in files:
            if codigo_busca in f.upper():
                return os.path.join(root, f)
    return None

SVG_SUZANO_OFICIAL = """<svg xmlns="http://www.w3.org/2000/svg" width="220" height="52" viewBox="0 0 120 28" fill="none">
<path d="M17.1158 9.99659C17.6873 7.6407 17.9414 5.28482 17.8778 2.86526V2.80159C22.0057 2.35588 26.6417 1.8465 29.3089 0C29.69 2.35588 30.071 6.23991 25.4351 12.7982C20.6721 19.4201 13.3054 23.0495 7.20884 24.9596C8.66948 23.8135 10.0031 22.4764 11.2732 21.0756C14.1945 17.7646 16.0997 14.008 17.1158 9.99659ZM15.4646 3.11995C9.74908 3.69301 5.7482 5.53951 3.20796 8.72314C0.4137 12.1615 -0.538889 17.4463 0.286688 24.6413C0.350194 25.3417 0.4137 25.851 0.477206 26.2331H0.540712C3.589 24.7686 15.5281 17.9557 15.4646 3.11995ZM31.0871 7.44969C30.198 9.86924 28.9914 12.1615 27.4672 14.2626C22.6408 20.9483 15.5281 24.5776 10.3206 26.4878C8.66948 27.1245 7.08183 7.5702 5.7482 27.9522C9.81259 28.1433 13.5594 27.7612 16.8617 26.8061C20.6086 25.7237 23.6569 23.8772 26.0701 21.394C29.5629 17.7646 31.3411 12.8619 31.0871 7.44969Z" fill="url(#paint0_radial_suzano)"/>
<path d="M47.2194 13.1188C46.3939 12.5457 44.7427 11.909 43.028 11.909C41.3134 11.909 40.1068 12.6731 40.1068 13.8192C40.1068 14.8379 41.0594 15.3473 42.5835 15.5383L44.1076 15.7294C46.8384 16.1114 48.4896 17.2575 48.4896 19.4224C48.4896 21.9056 46.2668 23.4974 42.774 23.4974C41.1864 23.4974 39.0272 23.1154 37.3125 21.8419L38.3286 20.1864C39.2812 20.9505 40.5513 21.5872 42.7105 21.5872C44.8697 21.5872 46.1398 20.8868 46.1398 19.6134C46.1398 18.6583 45.3143 18.0852 43.5361 17.8306L42.0119 17.6395C39.1542 17.2575 37.757 15.984 37.757 14.0102C37.757 11.527 39.7892 10.0625 42.9645 10.0625C44.9967 10.0625 46.8384 10.5719 48.1085 11.4633L47.2194 13.1188Z" fill="#FFFFFF"/>
<path d="M62.0152 17.7634C62.0152 21.6474 59.221 23.4302 56.1092 23.4302C52.9974 23.4302 50.2031 21.6474 50.2031 17.7634V10.25H52.4258V17.5723C52.4258 20.1829 54.0135 21.3927 56.1092 21.3927C58.2049 21.3927 59.7925 20.2466 59.7925 17.5723V10.25H62.0152V17.7634Z" fill="#FFFFFF"/>
<path d="M75.1564 11.8418L67.2181 21.138H75.2834V23.1755H64.2969V21.5837L72.3621 12.2875H64.5509V10.25H75.2199V11.8418H75.1564Z" fill="#FFFFFF"/>
<path d="M89.5071 23.1814H87.2844V21.4622C86.2683 22.672 84.7442 23.4361 82.712 23.4361C79.0921 23.4361 76.2344 20.6345 76.2344 16.6868C76.2344 12.7391 79.0286 9.9375 82.712 9.9375C84.7442 9.9375 86.2683 10.7016 87.2844 11.975V10.1922H89.5071V23.1814ZM78.5841 16.6868C78.5841 19.2337 80.3623 21.3349 83.0295 21.3349C85.6968 21.3349 87.4749 19.361 87.4749 16.6868C87.4749 14.0125 85.6333 12.0387 83.0295 12.0387C80.4258 12.0387 78.5841 14.1399 78.5841 16.6868Z" fill="#FFFFFF"/>
<path d="M104.375 15.1586V23.1814H102.152V15.859C102.152 13.5032 100.818 12.1024 98.5956 12.1024C96.4999 12.1024 94.7852 13.5032 94.7852 15.9227V23.1814H92.5625V10.2559H94.7852V11.9113C95.8648 10.4469 97.516 9.9375 98.9766 9.9375C102.215 9.9375 104.375 12.0387 104.375 15.1586Z" fill="#FFFFFF"/>
<path d="M119.998 16.6868C119.998 20.5071 117.076 23.4361 113.139 23.4361C109.202 23.4361 106.344 20.5071 106.344 16.6868C106.344 12.8664 109.265 9.9375 113.139 9.9375C117.013 9.9375 119.998 12.8664 119.998 16.6868ZM108.63 16.6868C108.63 19.4247 110.599 21.3349 113.202 21.3349C115.806 21.3349 117.775 19.361 117.775 16.6868C117.775 14.0125 115.806 12.0387 113.202 12.0387C110.599 12.0387 108.63 14.0125 108.63 16.6868Z" fill="#FFFFFF"/>
<defs>
<radialGradient id="paint0_radial_suzano" cx="0" cy="0" r="1" gradientUnits="userSpaceOnUse" gradientTransform="translate(-8.33107 14.1226) scale(45.6291 45.7487)">
<stop stop-color="#A4E03D"/>
<stop offset="0.5" stop-color="#00CC66"/>
<stop offset="1" stop-color="#008461"/>
</radialGradient>
</defs>
</svg>"""

b64_suzano_logo = base64.b64encode(SVG_SUZANO_OFICIAL.encode('utf-8')).decode('utf-8')
URL_LOGO_TAGLINE_RODAPE = "https://cdn.prod.website-files.com/69cb021343dbf125cabfa994/69cb021343dbf125cabfa9df_logo-tagline.svg"

st.markdown("""
<style>
    div.stButton > button:first-child, div.stDownloadButton > button:first-child {
        background-color: #03267F !important;
        color: #FFFFFF !important;
        border-radius: 8px;
        border: 1px solid rgba(255,255,255,0.1);
        padding: 0.6rem 1.2rem;
        font-weight: 700;
        transition: all 0.3s ease;
    }
    div.stButton > button:first-child:hover, div.stDownloadButton > button:first-child:hover {
        background-color: #00C85A !important;
        color: #FFFFFF !important;
        border-color: #00C85A !important;
        transform: translateY(-1px);
    }
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
</style>
""", unsafe_allow_html=True)

banner_html = f"""
<div style="background: linear-gradient(95deg, #03267F 0%, #003882 60%, #004B9B 100%); padding: 26px 35px; border-radius: 12px 12px 0px 0px; display: flex; align-items: center; justify-content: space-between; flex-wrap: wrap; gap: 20px; box-shadow: 0 6px 16px rgba(0,0,0,0.35);">
    <div style="display: flex; align-items: center; gap: 24px;">
        <img src="data:image/svg+xml;base64,{b64_suzano_logo}" style="height: 48px; display: block;" alt="Logo Suzano Oficial">
        <div style="height: 38px; width: 2px; background-color: rgba(255,255,255,0.2);"></div>
        <div>
            <div style="color: #A4E03D; font-size: 13px; font-weight: 800; text-transform: uppercase; letter-spacing: 2px;">Trade Marketing • Bens de Consumo</div>
            <div style="color: #FFFFFF; font-size: 20px; font-weight: 600; margin-top: 2px;">Gerador Oficial de Guias de Execução Checkout</div>
        </div>
    </div>
    <div style="text-align: right;">
        <span style="background-color: rgba(0, 200, 90, 0.2); border: 1px solid #00C85A; color: #A4E03D; padding: 6px 14px; border-radius: 20px; font-size: 12px; font-weight: 600;">Portfólio Higiene 2026</span>
    </div>
</div>
<div style="display: flex; height: 7px; border-radius: 0px 0px 12px 12px; margin-bottom: 30px; overflow: hidden;">
    <div style="flex: 1; background-color: #00C85A;" title="Suzano"></div>
    <div style="flex: 1; background-color: #D4AF37;" title="Neve"></div>
    <div style="flex: 1; background-color: #662D91;" title="Mimmo"></div>
    <div style="flex: 1; background-color: #009CDE;" title="Scott / Duramax"></div>
    <div style="flex: 1; background-color: #0055A4;" title="Kleenex"></div>
</div>
"""
st.markdown(banner_html, unsafe_allow_html=True)

CATALOGO_PADRAO = [
    ("20104429", "7891172151125", "GUARDANAPO SCOTT GRAND HOTEL FD 9X9X50P", "Guardanapo", "Scott"),
    ("20104430", "7891172151323", "GUARDANAPO SCOTT GRAND HOTEL FD 8X9X50G", "Guardanapo", "Scott"),
    ("20104426", "7891172145124", "GUARD. SCOTT DIA A DIA FS 6X12X50 P", "Guardanapo", "Scott"),
    ("20104427", "7891172145322", "GUARD. SCOTT DIA A DIA FS 6X12X50 G", "Guardanapo", "Scott"),
    ("20104420/20104421", "7891172161162 / 7891172172106", "LENÇO DE PAPEL KLEENEX 50 ou 50/60 LENÇOS", "Lenços Secos", "Kleenex"),
    ("20104309", "7891172523380", "LENÇO DE PAPEL KLEENEX L4P3", "Lenços Secos", "Kleenex"),
    ("20104422", "7891172523366", "LENÇO DE PAPEL KLEENEX 100/80 LENÇOS", "Lenços Secos", "Kleenex"),
    ("20104423", "7891172523373", "LENÇO DE PAPEL KLEENEX 150 LENÇOS", "Lenços Secos", "Kleenex"),
    ("20104310", "7891172171666", "LENÇO DE PAPEL KLEENEX 12 PACOTES (BOLSO)", "Lenços Secos", "Kleenex"),
    ("90005190", "7891172523908", "NEVE WIPES CARE 48 LENÇOS", "Wipes", "Neve"),
    ("90005191", "7891172523915", "NEVE WIPES CARE LEVE 4 PAGUE 3", "Wipes", "Neve"),
    ("90004710", "7898962794241", "LENÇO UMEDECIDO MIMMO 40 LENÇOS", "Wipes", "Mimmo"),
    ("90005082", "7891172523816", "NEVE WIPES ON THE GO", "Wipes", "Neve"),
    ("20104313", "7891172523328", "TC SCOTT DURAMAX 1R BRANCO", "Pano Reutilizável", "Scott"),
    ("20104312", "7891172523311", "TC SCOTT DURAMAX 1R AZUL", "Pano Reutilizável", "Scott"),
    ("20104311", "7891172523359", "PANO SCOTT DURAMAX 48X01X12 ANTIBAC", "Pano Reutilizável", "Scott"),
    ("20104424", "7891172523335", "PANO SCOTT DURAMAX 58X02X12 BRANCO (2R)", "Pano Reutilizável", "Scott"),
    ("20081465", "7898962794166", "MIMMO FOLHA DUPLA 12 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Mimmo"),
    ("20104405", "7891172523434", "NEVE FOLHA DUPLA 12 ROLOS 20M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20104408", "7891172421174", "NEVE FOLHA DUPLA 4 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20104409", "7891172523410", "NEVE FOLHA DUPLA 12 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20081467", "7898962794173", "MIMMO FOLHA DUPLA 16 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Mimmo"),
    ("20104410", "7891172432026", "NEVE FOLHA DUPLA 16 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20081471", "7898962794142", "MIMMO FOLHA DUPLA 4 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Mimmo"),
    ("20104411", "7891172523403", "NEVE FD L18P16 30M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20104412", "7891172433078", "NEVE FOLHA DUPLA 24 ROLOS 30M", "Papel Higiênico - Folha Dupla", "Neve"),
    ("20104415", "7891172523472", "NEVE FOLHA TRIPLA 12 ROLOS 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20097397", "7898962794265", "MIMMO FOLHA TRIPLA 12 ROLOS 20M", "Papel Higiênico - Folha Tripla", "Mimmo"),
    ("20104414", "7891172422607", "NEVE FOLHA TRIPLA 4 ROLOS 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20104416", "7891172432804", "NEVE FOLHA TRIPLA 16 ROLOS 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20104417", "7891172523465", "NEVE FT L18P16 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20104418", "7891172433085", "NEVE FOLHA TRIPLA 24 ROLOS 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20109736", "7891172523847", "NEVE FOLHA TRIPLA TOQUE DE ONDAS 6R 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20109735", "7891172523861", "NEVE FOLHA TRIPLA TOQUE DE ONDAS 12R 20M", "Papel Higiênico - Folha Tripla", "Neve"),
    ("20106705", "7891172523717", "NEVE FOLHA QUADRUPLA 4R 20M", "Papel Higiênico - Folha Quádrupla", "Neve"),
    ("20106704", "7891172523724", "NEVE FOLHA QUADRUPLA 12R 20M", "Papel Higiênico - Folha Quádrupla", "Neve"),
]

def gerar_excel_template():
    linhas = []
    for cluster in ["CK 1 A 4 CK", "CK 05 A 09 CK", "CK 10 A 19 CK"]:
        for mat, ean, prod, cat, marca in CATALOGO_PADRAO:
            if cluster in ["CK 1 A 4 CK", "CK 05 A 09 CK"] and mat in ["20104411", "20104412", "20104417", "20104418"]:
                status_padrao = "EXCLUIR"
            else:
                status_padrao = "COMPLEMENTAR"
            linhas.append({
                "CLUSTER": cluster, "MATERIAL": mat, "EAN": ean,
                "PRODUTO": prod, "CATEGORIA": cat, "MARCA": marca, "PRIORIDADE": status_padrao
            })
    df_temp = pd.DataFrame(linhas)
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_temp.to_excel(writer, index=False, sheet_name="Mix_Regional")
    return output.getvalue()

col_esquerda, space, col_direita = st.columns([1.05, 0.08, 1.25])

with col_esquerda:
    st.markdown("### 📥 1. Etapa de Planeamento")
    st.write("Transfira o ficheiro Excel oficial com o portfólio completo da Suzano. Edite a coluna **PRIORIDADE** conforme a estratégia do seu cliente.")
    st.download_button(
        label="📄 Descarregar Excel Template (.xlsx)",
        data=gerar_excel_template(),
        file_name="Template_Mix_Suzano_2026.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("### ⚙️ 2. Gerar Manual de Execução")
    st.write("Preencha os dados da conta e faça o upload da folha de cálculo configurada.")
    
    titulo_regional = st.text_input("📝 Nome do Distribuidor / Conta:", placeholder="Ex: DESTRO - PR / SC")
    arquivo_up = st.file_uploader("📂 Suba o Excel Template preenchido (.xlsx)", type=["xlsx"])
    
    gerar = st.button("🚀 Processar e Gerar PowerPoint", type="primary", use_container_width=True)

with col_direita:
    st.markdown("### 🏢 As Nossas Marcas & Destaques")
    st.write("Catálogo Oficial - Segmento de Higiene e Cuidados Pessoais")
    
    logo_cols = st.columns(5)
    marcas_vitrine = ["NEVE", "MIMMO", "SCOTT", "KLEENEX", "DURAMAX"]
    for i, m_nome in enumerate(marcas_vitrine):
        l_path = achar_logo(m_nome)
        if l_path:
            logo_cols[i].image(l_path, use_container_width=True)
        else:
            logo_cols[i].markdown(f"<div style='text-align:center; padding:12px; font-weight:bold; color:#fff; background:rgba(255,255,255,0.08); border: 1px solid rgba(255,255,255,0.15); border-radius:6px;'>{m_nome}</div>", unsafe_allow_html=True)
            
    st.markdown("<hr style='margin: 15px 0px; border-color: #444444;'>", unsafe_allow_html=True)
    
    prod_cols_1 = st.columns(4)
    prods_linha_1 = ["20104405", "20106704", "20081465", "90005190"]
    for i, p_mat in enumerate(prods_linha_1):
        p_path = achar_imagem(p_mat)
        if p_path: prod_cols_1[i].image(p_path, use_container_width=True)

    st.markdown("<div style='margin-top: 10px;'></div>", unsafe_allow_html=True)

    prod_cols_2 = st.columns(4)
    prods_linha_2 = ["20104429", "20104420", "20104312", "20104424"]
    for i, p_mat in enumerate(prods_linha_2):
        p_path = achar_imagem(p_mat)
        if p_path: prod_cols_2[i].image(p_path, use_container_width=True)

footer_html = f"""
<div style="margin-top: 50px; padding-top: 20px; border-top: 1px solid rgba(255,255,255,0.1); display: flex; align-items: center; justify-content: space-between; flex-wrap: wrap;">
    <img src="{URL_LOGO_TAGLINE_RODAPE}" style="height: 32px; opacity: 0.9;" alt="Suzano - Nós plantamos o futuro">
    <span style="color: #888888; font-size: 12px;">® Suzano S/A - Todos os direitos reservados • Ferramenta Interna de Trade Marketing</span>
</div>
"""
st.markdown(footer_html, unsafe_allow_html=True)

if gerar and arquivo_up:
    if not titulo_regional:
        st.warning("Por favor, preencha o nome do Distribuidor/Conta antes de gerar o manual.")
    else:
        with st.spinner("A ligar ao banco de imagens Suzano e a renderizar os diapositivos..."):
            
            df = pd.read_excel(arquivo_up)
            df.columns = [str(c).strip().upper() for c in df.columns]
            df = df[~df['PRIORIDADE'].astype(str).str.upper().str.contains('EXCLUIR|NAO|NÃO|FORA')]
            
            def extrair_num(txt):
                nums = re.findall(r'\d+', str(txt))
                return int(nums[0]) if nums else 99

            df['Cluster_Num'] = df['CLUSTER'].apply(extrair_num)
            ordem_cat = {
                'Guardanapo': 1, 'Lenços Secos': 2, 'Wipes': 3, 'Pano Reutilizável': 4,
                'Papel Higiênico - Folha Dupla': 5, 'Papel Higiênico - Folha Tripla': 6, 'Papel Higiênico - Folha Quádrupla': 7
            }
            ordem_prio = {'PRIORITARIO E HERO': 1, 'PRIORITARIO': 2, 'COMPLEMENTAR': 3}
            
            df['cat_ord'] = df['CATEGORIA'].map(ordem_cat).fillna(99)
            df['prio_ord'] = df['PRIORIDADE'].astype(str).str.upper().str.strip().map(ordem_prio).fillna(99)
            df = df.sort_values(by=['Cluster_Num', 'cat_ord', 'prio_ord']).reset_index(drop=True)

            prs_out = Presentation()
            prs_out.slide_width, prs_out.slide_height = Inches(13.333), Inches(7.5)
            blank_layout = prs_out.slide_layouts[6]

            slide_capa = prs_out.slides.add_slide(blank_layout)
            if os.path.exists(ARQUIVO_CAPA):
                prs_base = Presentation(ARQUIVO_CAPA)
                for sh in prs_base.slides[0].shapes:
                    if sh.shape_type == 13:
                        slide_capa.shapes.add_picture(BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
            else:
                bg_c = slide_capa.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
                bg_c.fill.solid()
                bg_c.fill.fore_color.rgb = RGBColor(0, 56, 130)

            tb_capa = slide_capa.shapes.add_textbox(Inches(2.5), Inches(2.8), Inches(8.5), Inches(2.0))
            tb_capa.text_frame.paragraphs[0].text = f"GUIA DE VENDAS CHECKOUT\n{titulo_regional.upper()}"
            tb_capa.text_frame.paragraphs[0].font.size = Pt(36)
            tb_capa.text_frame.paragraphs[0].font.bold = True
            tb_capa.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
            tb_capa.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

            CARDS_POS = [
                {"x": Inches(0.35), "y": Inches(0.80)}, {"x": Inches(6.80), "y": Inches(0.80)},
                {"x": Inches(0.35), "y": Inches(3.95)}, {"x": Inches(6.80), "y": Inches(3.95)}
            ]

            logo_suzano = achar_logo("SUZANO")
            logo_scott = achar_logo("SCOTT")
            logo_kleenex = achar_logo("KLEENEX")
            logo_duramax = achar_logo("DURAMAX")
            logo_neve = achar_logo("NEVE")

            for cluster, group in df.groupby("CLUSTER", sort=False):
                for subcat, grp_sub in group.groupby("CATEGORIA", sort=False):
                    itens = grp_sub.to_dict("records")
                    total_items = len(itens)
                    for i in range(0, total_items, 4):
                        chunk = itens[i:i+4]
                        slide = prs_out.slides.add_slide(blank_layout)
                        
                        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
                        bg.fill.solid()
                        bg.fill.fore_color.rgb = RGBColor(242, 244, 248)
                        bg.line.fill.background()
                        
                        top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.62))
                        top.fill.solid()
                        top.fill.fore_color.rgb = RGBColor(0, 56, 130)
                        top.line.fill.background()
                        
                        tb_cl = slide.shapes.add_textbox(Inches(0.35), Inches(0.08), Inches(4.5), Inches(0.45))
                        tb_cl.text_frame.paragraphs[0].text = str(cluster).upper()
                        tb_cl.text_frame.paragraphs[0].font.size = Pt(22)
                        tb_cl.text_frame.paragraphs[0].font.bold = True
                        tb_cl.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                        
                        tb_sub = slide.shapes.add_textbox(Inches(4.5), Inches(0.12), Inches(4.5), Inches(0.40))
                        tb_sub.text_frame.paragraphs[0].text = str(subcat).split(" - ")[-1].upper()
                        tb_sub.text_frame.paragraphs[0].font.size = Pt(18)
                        tb_sub.text_frame.paragraphs[0].font.bold = True
                        tb_sub.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                        tb_sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

                        tb_it = slide.shapes.add_textbox(Inches(10.0), Inches(0.14), Inches(3.0), Inches(0.40))
                        tb_it.text_frame.paragraphs[0].text = f"Itens {i+1}-{min(i+4, total_items)}"
                        tb_it.text_frame.paragraphs[0].font.size = Pt(15)
                        tb_it.text_frame.paragraphs[0].font.bold = True
                        tb_it.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                        tb_it.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                        
                        bot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.85), Inches(13.333), Inches(0.65))
                        bot.fill.solid()
                        bot.fill.fore_color.rgb = RGBColor(0, 70, 155)
                        bot.line.fill.background()
                        
                        if logo_suzano: slide.shapes.add_picture(logo_suzano, Inches(11.6), Inches(6.88), height=Inches(0.58))
                        
                        if "Guardanapo" in subcat and logo_scott:
                            slide.shapes.add_picture(logo_scott, Inches(0.35), Inches(6.86), height=Inches(0.60))
                        elif "Lenços" in subcat and logo_kleenex:
                            slide.shapes.add_picture(logo_kleenex, Inches(0.35), Inches(6.88), height=Inches(0.55))
                        elif "Pano" in subcat and logo_duramax:
                            slide.shapes.add_picture(logo_duramax, Inches(0.35), Inches(6.85), height=Inches(0.62))
                        elif logo_neve:
                            slide.shapes.add_picture(logo_neve, Inches(0.35), Inches(6.86), height=Inches(0.58))
                        
                        for j, item in enumerate(chunk):
                            pos = CARDS_POS[j]
                            cx, cy = pos["x"], pos["y"]
                            imp_num = i + j + 1
                            
                            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(6.15), Inches(2.95))
                            card.fill.solid()
                            card.fill.fore_color.rgb = RGBColor(255, 255, 255)
                            card.line.color.rgb = RGBColor(200, 210, 225)
                            card.line.width = Pt(1.5)
                            
                            cbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(6.15), Inches(0.42))
                            cbar.fill.solid()
                            cbar.fill.fore_color.rgb = RGBColor(0, 75, 160)
                            cbar.line.fill.background()
                            
                            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.15), cy + Inches(0.04), Inches(2.2), Inches(0.34))
                            badge.fill.solid()
                            badge.fill.fore_color.rgb = RGBColor(10, 100, 200)
                            badge.line.color.rgb = RGBColor(120, 180, 255)
                            
                            p_b = badge.text_frame.paragraphs[0]
                            p_b.text = str(item['CATEGORIA']).split(' - ')[0].upper()
                            p_b.font.size = Pt(11)
                            p_b.font.bold = True
                            p_b.font.color.rgb = RGBColor(255, 255, 255)
                            p_b.alignment = PP_ALIGN.CENTER
                            
                            mat_code = str(item['MATERIAL'])
                            img_path = achar_imagem(mat_code)
                            
                            if img_path:
                                slide.shapes.add_picture(img_path, cx + Inches(0.15), cy + Inches(0.52), width=Inches(1.85))
                            else:
                                ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.15), cy + Inches(0.55), Inches(1.85), Inches(1.8))
                                ph.fill.solid()
                                ph.fill.fore_color.rgb = RGBColor(240, 240, 240)
                                ph.line.color.rgb = RGBColor(180, 180, 180)
                                
                                p_ph = ph.text_frame.paragraphs[0]
                                p_ph.text = f"Inserir Foto\n(Mat: {mat_code})"
                                p_ph.font.size = Pt(9)
                                p_ph.font.bold = True
                                p_ph.font.color.rgb = RGBColor(100, 100, 100)
                                p_ph.alignment = PP_ALIGN.CENTER
                            
                            tb_info = slide.shapes.add_textbox(cx + Inches(2.10), cy + Inches(0.48), Inches(3.95), Inches(1.90))
                            tf_i = tb_info.text_frame
                            tf_i.word_wrap = True
                            
                            p_n = tf_i.paragraphs[0]
                            p_n.text = str(item['PRODUTO'])
                            p_n.font.size = Pt(12)
                            p_n.font.bold = True
                            p_n.font.color.rgb = RGBColor(30, 30, 30)
                            
                            for lbl in [f"Categoria: {str(item['CATEGORIA']).split(' - ')[0]}", f"EAN: {item['EAN']}", f"Material: {mat_code}"]:
                                p_sub = tf_i.add_paragraph()
                                p_sub.text = lbl
                                p_sub.font.size = Pt(10.5)
                                p_sub.font.color.rgb = RGBColor(80, 80, 80)
                            
                            status_str = str(item['PRIORIDADE']).upper().strip()
                            tb_prio = slide.shapes.add_textbox(cx + Inches(0.15), cy + Inches(2.48), Inches(3.8), Inches(0.38))
                            
                            p_p = tb_prio.text_frame.paragraphs[0]
                            p_p.text = f"Prioridade: {status_str}"
                            p_p.font.size = Pt(11)
                            p_p.font.bold = True
                            
                            if "HERO" in status_str:
                                p_p.font.color.rgb = RGBColor(0, 130, 60)
                            elif "PRIORITARIO" in status_str:
                                p_p.font.color.rgb = RGBColor(0, 90, 180)
                            else:
                                p_p.font.color.rgb = RGBColor(90, 90, 90)
                            
                            tb_imp = slide.shapes.add_textbox(cx + Inches(4.10), cy + Inches(2.48), Inches(1.90), Inches(0.38))
                            p_imp = tb_imp.text_frame.paragraphs[0]
                            p_imp.text = f"Importância: {imp_num}"
                            p_imp.font.size = Pt(11)
                            p_imp.font.bold = True
                            p_imp.font.color.rgb = RGBColor(0, 75, 160)
                            p_imp.alignment = PP_ALIGN.RIGHT

            ppt_bytes = BytesIO()
            prs_out.save(ppt_bytes)
            st.success("✅ Apresentação Finalizada com Sucesso! Clique abaixo para baixar.")
            st.download_button(
                label="📥 Baixar Guia de Vendas (.pptx)",
                data=ppt_bytes.getvalue(),
                file_name=f"Guia_Mix_{titulo_regional.replace(' ', '_')}.pptx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.presentation"
            )
